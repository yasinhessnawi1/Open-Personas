"""The parsers' own bounds: what a file can make them do AFTER the archive guard passes it.

The guard bounds bytes; these bound the parsers' logic, which amplifies on its own: an
xlsx padded to its declared dimension, a docx cell repeated across its span, a PDF whose
pages all share one stream. Every bound is crossed in its real unit, one step at the
limit (passes) and one past it (bounded), and each security review regression file is
rebuilt here.

A test whose hostile file would allocate without limit if its bound regressed parses
in a memory-limited child (``bounded_parse``), so a regression fails one test instead of
killing the test run.
"""

# ruff: noqa: E501 (OOXML content-type and relationship URIs are long literals)
from __future__ import annotations

import io
import zipfile
from typing import TYPE_CHECKING

import docx
import pytest
from persona.documents.archive_guard import check_ooxml_archive
from persona.documents.errors import CorruptDocumentError
from persona.documents.parsers import (
    FLAG_PARTIAL_EXTRACTION,
    FLAG_ROW_CAP_TRUNCATED,
    parse_document,
)
from persona.documents.parsers import docx as docx_parser
from persona.documents.parsers import pdf as pdf_parser
from persona.documents.parsers import xlsx as xlsx_parser

from tests.unit.documents._bounded_parse import bounded_parse

if TYPE_CHECKING:
    from pathlib import Path


def _live(name: str) -> object:
    """The parser module ``parse_document`` will call NOW.

    Other tests reload the parser modules (to simulate a missing library), so a module
    object imported at the top of this file can be stale by the time a test runs.
    """
    import importlib

    return importlib.import_module(f"persona.documents.parsers.{name}")


# --------------------------------------------------------------------------- builders

_SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _xlsx(
    path: Path,
    *,
    dimension: str | None,
    rows_xml: str,
    shared_strings: list[str] | None = None,
) -> Path:
    """A minimal valid xlsx whose single sheet is exactly ``rows_xml``."""
    ct = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    workbook = (
        f'<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="{_SHEET_NS}" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="S" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    workbook_rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
        "</Relationships>"
    )
    dim = f'<dimension ref="{dimension}"/>' if dimension else ""
    sheet = (
        f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{_SHEET_NS}">'
        f"{dim}<sheetData>{rows_xml}</sheetData></worksheet>"
    )
    strings = shared_strings or []
    sst = (
        f'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="{_SHEET_NS}" '
        f'count="{len(strings)}" uniqueCount="{len(strings)}">'
        + "".join(f"<si><t>{s}</t></si>" for s in strings)
        + "</sst>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", ct)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        archive.writestr("xl/worksheets/sheet1.xml", sheet)
        archive.writestr("xl/sharedStrings.xml", sst)
    return path


def _column(index: int) -> str:
    """1-based column number to its letters (1 -> A, 27 -> AA)."""
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _inline(ref: str, text: str) -> str:
    return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'


def _row(number: int, cells: str) -> str:
    return f'<row r="{number}">{cells}</row>'


_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _docx(path: Path, body_xml: str) -> Path:
    """A real python-docx package whose document body is exactly ``body_xml``."""
    buffer = io.BytesIO()
    seed = docx.Document()
    seed.add_paragraph("seed")
    seed.save(buffer)
    document_xml = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f"<w:document {_W}><w:body>{body_xml}</w:body></w:document>"
    ).encode()
    with (
        zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source,
        zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out,
    ):
        for info in source.infolist():
            data = (
                document_xml if info.filename == "word/document.xml" else source.read(info.filename)
            )
            out.writestr(info.filename, data)
    return path


_PARAGRAPH = "<w:p><w:r><w:t>hello world, this is a real paragraph</w:t></w:r></w:p>"


def _pdf(path: Path, *, pages: int, declared: int | None = None, text: str = "hello") -> Path:
    """A PDF whose ``pages`` pages all point at ONE shared content stream.

    ``declared`` sets the page tree's ``/Count`` independently, so it can lie.
    """
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # page tree, filled in below
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
    ]
    first_page = len(objects) + 1
    kids = " ".join(f"{first_page + i} 0 R" for i in range(pages))
    count = pages if declared is None else declared
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode()
    for _ in range(pages):
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents 4 0 R >>"
        )
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    path.write_bytes(bytes(out))
    return path


# --------------------------------------------------------------------------- xlsx


def test_the_reviewers_gap_workbook_is_bounded(tmp_path: Path) -> None:
    """``A1:XFD1000`` with two cells: 383 MiB and 49 MB of text before the fix."""
    path = _xlsx(
        tmp_path / "gap.xlsx",
        dimension="A1:XFD1000",
        rows_xml=_row(1, _inline("A1", "hi")) + _row(1000, _inline("A1000", "x")),
    )
    check_ooxml_archive(path, document_format="xlsx")  # precondition: the guard passes it
    result = bounded_parse(path).result
    assert len(result.full_text) < 10_000
    assert FLAG_PARTIAL_EXTRACTION in result.flags  # columns past the cap were dropped
    assert "only the first 256 columns" in result.full_text


def test_the_declared_width_is_read_at_most_the_column_cap_per_row(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The WORK, in cells: openpyxl pads every row to the width it is asked for.

    Asked for the declared 16,384 columns it pads each gap row to 16,384 cells; asked for
    the capped width it pads to 256. The text alone cannot show the difference, because
    the rows are cut afterwards either way.
    """
    module = _live("xlsx")
    spent: list[int] = []
    real_budget = module._Budget

    class _Recording(real_budget):  # type: ignore[misc,valid-type]
        def __init__(self) -> None:
            super().__init__()
            spent.append(0)
            self._slot = len(spent) - 1

        def __setattr__(self, name: str, value: object) -> None:
            super().__setattr__(name, value)
            if name == "cells" and hasattr(self, "_slot"):
                spent[self._slot] = int(value)  # type: ignore[call-overload]

    monkeypatch.setattr(module, "_Budget", _Recording)
    path = _xlsx(
        tmp_path / "gap.xlsx",
        dimension="A1:XFD1000",
        rows_xml=_row(1, _inline("A1", "hi")) + _row(1000, _inline("A1000", "x")),
    )
    parse_document(path)
    assert spent, "precondition: the budget was used"
    assert spent[0] == 1000 * module.MAX_SHEET_COLUMNS


def test_a_wide_row_in_a_sheet_without_a_dimension_is_cut(tmp_path: Path) -> None:
    """No ``<dimension>``: openpyxl sizes each row to its last cell; the row is cut here."""
    columns = xlsx_parser.MAX_SHEET_COLUMNS + 1
    cells = "".join(_inline(f"{_column(c)}1", f"h{c}") for c in range(1, columns + 1))
    cells2 = "".join(_inline(f"{_column(c)}2", f"v{c}") for c in range(1, columns + 1))
    path = _xlsx(tmp_path / "nodim.xlsx", dimension=None, rows_xml=_row(1, cells) + _row(2, cells2))
    result = parse_document(path)
    assert FLAG_PARTIAL_EXTRACTION in result.flags
    assert f"v{columns - 1}" in result.full_text
    assert f"v{columns}" not in result.full_text


def test_the_reviewers_maximal_gap_workbook_is_bounded(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path / "max.xlsx",
        dimension="A1:XFD1048576",
        rows_xml=_row(1, _inline("A1", "hi")) + _row(1048576, _inline("A1048576", "x")),
    )
    result = bounded_parse(path).result
    assert len(result.full_text) < 20_000
    # The padded gap rows exhaust the cell budget long before the row limit.
    assert "reading stopped after row" in result.full_text
    assert FLAG_PARTIAL_EXTRACTION in result.flags


def _wide(tmp_path: Path, columns: int) -> Path:
    cells = "".join(_inline(f"{_column(c)}1", f"h{c}") for c in range(1, columns + 1))
    cells2 = "".join(_inline(f"{_column(c)}2", f"v{c}") for c in range(1, columns + 1))
    return _xlsx(
        tmp_path / f"w{columns}.xlsx",
        dimension=f"A1:{_column(columns)}2",
        rows_xml=_row(1, cells) + _row(2, cells2),
    )


def test_a_sheet_exactly_at_the_column_cap_is_read_whole(tmp_path: Path) -> None:
    columns = xlsx_parser.MAX_SHEET_COLUMNS
    result = parse_document(_wide(tmp_path, columns))
    assert FLAG_PARTIAL_EXTRACTION not in result.flags
    assert f"v{columns}" in result.full_text


def test_a_sheet_one_column_past_the_cap_is_cut_and_flagged(tmp_path: Path) -> None:
    columns = xlsx_parser.MAX_SHEET_COLUMNS + 1
    result = parse_document(_wide(tmp_path, columns))
    assert FLAG_PARTIAL_EXTRACTION in result.flags
    assert f"v{columns - 1}" in result.full_text
    assert f"v{columns}" not in result.full_text


def _tall(tmp_path: Path, last_row: int) -> Path:
    rows = _row(1, _inline("A1", "h")) + _row(last_row, _inline(f"A{last_row}", "end"))
    return _xlsx(tmp_path / f"t{last_row}.xlsx", dimension=f"A1:A{last_row}", rows_xml=rows)


def test_a_sheet_exactly_at_the_row_limit_is_read_to_its_end(tmp_path: Path) -> None:
    result = parse_document(_tall(tmp_path, xlsx_parser.MAX_SHEET_ROWS))
    assert FLAG_PARTIAL_EXTRACTION not in result.flags
    assert "end" in result.full_text


def test_a_sheet_one_row_past_the_limit_stops_and_says_so(tmp_path: Path) -> None:
    result = parse_document(_tall(tmp_path, xlsx_parser.MAX_SHEET_ROWS + 1))
    assert FLAG_PARTIAL_EXTRACTION in result.flags
    assert "end" not in result.full_text
    assert f"reading stopped after row {xlsx_parser.MAX_SHEET_ROWS}" in result.full_text


def _cells(tmp_path: Path, rows: int, width: int) -> Path:
    """``rows`` x ``width`` cells read (padding counts), the last row holding a value."""
    body = _row(1, _inline("A1", "h")) + _row(rows, _inline(f"A{rows}", "end"))
    return _xlsx(tmp_path / f"c{rows}.xlsx", dimension=f"A1:{_column(width)}{rows}", rows_xml=body)


def test_a_workbook_exactly_at_the_cell_budget_is_read_whole(tmp_path: Path) -> None:
    width = 250
    rows = xlsx_parser.MAX_WORKBOOK_CELLS // width
    assert rows * width == xlsx_parser.MAX_WORKBOOK_CELLS
    result = parse_document(_cells(tmp_path, rows, width))
    assert FLAG_PARTIAL_EXTRACTION not in result.flags
    assert "end" in result.full_text


def test_a_workbook_one_row_past_the_cell_budget_stops(tmp_path: Path) -> None:
    width = 250
    rows = xlsx_parser.MAX_WORKBOOK_CELLS // width + 1
    result = parse_document(_cells(tmp_path, rows, width))
    assert FLAG_PARTIAL_EXTRACTION in result.flags
    assert "end" not in result.full_text


def test_one_shared_string_referenced_everywhere_is_cut_per_cell(tmp_path: Path) -> None:
    long_string = "z" * 100_000
    cells = "".join(f'<c r="{_column(c)}1" t="s"><v>0</v></c>' for c in range(1, 41))
    rows = "".join(_row(r, cells.replace('1"', f'{r}"')) for r in range(1, 51))
    path = _xlsx(
        tmp_path / "s.xlsx", dimension="A1:AN50", rows_xml=rows, shared_strings=[long_string]
    )
    result = bounded_parse(path).result
    assert len(result.full_text) <= 50 * 40 * (xlsx_parser.MAX_CELL_CHARS + 3) + 1_000
    assert "z" * xlsx_parser.MAX_CELL_CHARS not in result.full_text


def _text_heavy(tmp_path: Path, rows: int) -> Path:
    cell = "y" * (xlsx_parser.MAX_CELL_CHARS - 1)
    cells = "".join(f'<c r="{_column(c)}1" t="s"><v>0</v></c>' for c in range(1, 65))
    body = "".join(_row(r, cells.replace('1"', f'{r}"')) for r in range(1, rows + 1))
    return _xlsx(
        tmp_path / f"x{rows}.xlsx", dimension=f"A1:BL{rows}", rows_xml=body, shared_strings=[cell]
    )


def test_the_workbook_text_budget_stops_keeping_rows(tmp_path: Path) -> None:
    # 64 cells of 255 characters cost 64 * (255 + 3) per row: about 242 rows reach 4 M.
    per_row = 64 * (xlsx_parser.MAX_CELL_CHARS - 1 + 3)
    fits = xlsx_parser.MAX_WORKBOOK_TEXT_CHARS // per_row
    inside = parse_document(_text_heavy(tmp_path, fits - 1))
    assert FLAG_PARTIAL_EXTRACTION not in inside.flags
    over = parse_document(_text_heavy(tmp_path, fits + 2))
    assert FLAG_PARTIAL_EXTRACTION in over.flags
    assert len(over.full_text) < xlsx_parser.MAX_WORKBOOK_TEXT_CHARS


def test_a_genuine_ledger_still_parses_with_its_real_tail(tmp_path: Path) -> None:
    """Valid input passes: a real 20,000-row workbook, the summary carries its last row."""
    import openpyxl

    workbook = openpyxl.Workbook(write_only=True)
    sheet = workbook.create_sheet("Ledger")
    sheet.append(["id", "amount", "note"])
    for row in range(20_000):
        sheet.append([row, row * 1.5, f"entry {row}"])
    path = tmp_path / "ledger.xlsx"
    workbook.save(path)
    result = parse_document(path)
    assert FLAG_ROW_CAP_TRUNCATED in result.flags
    assert FLAG_PARTIAL_EXTRACTION not in result.flags
    assert "20000 data rows" in result.full_text
    assert "entry 19999" in result.full_text
    assert "entry 0" in result.full_text


def test_a_memory_error_in_an_xlsx_is_never_relabelled(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _boom(*_args: object, **_kwargs: object) -> object:
        raise MemoryError

    monkeypatch.setattr(_live("xlsx"), "_read_sheet", _boom)
    path = _xlsx(tmp_path / "m.xlsx", dimension="A1:A1", rows_xml=_row(1, _inline("A1", "x")))
    with pytest.raises(MemoryError):
        parse_document(path)


def test_a_recursion_error_in_an_xlsx_is_a_clean_refusal(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _deep(*_args: object, **_kwargs: object) -> object:
        raise RecursionError

    monkeypatch.setattr(_live("xlsx"), "_read_sheet", _deep)
    path = _xlsx(tmp_path / "r.xlsx", dimension="A1:A1", rows_xml=_row(1, _inline("A1", "x")))
    with pytest.raises(CorruptDocumentError) as caught:
        parse_document(path)
    assert caught.value.context["reason"] == "too_complex"


# --------------------------------------------------------------------------- docx


def _span_table(span: int) -> str:
    return (
        _PARAGRAPH + "<w:tbl><w:tr><w:tc><w:tcPr>"
        f'<w:gridSpan w:val="{span}"/></w:tcPr><w:p><w:r><w:t>x</w:t></w:r></w:p>'
        "</w:tc></w:tr></w:tbl>"
    )


@pytest.mark.parametrize("span", [100_000, 2_147_483_647])
def test_the_reviewers_grid_span_produces_its_text_once(tmp_path: Path, span: int) -> None:
    result = bounded_parse(_docx(tmp_path / f"g{span}.docx", _span_table(span))).result
    assert len(result.full_text) < 200


def test_a_spanned_cell_appears_once(tmp_path: Path) -> None:
    """The own walk reads each ``w:tc`` once; ``row.cells`` repeated it per spanned column."""
    body = (
        _PARAGRAPH + "<w:tbl><w:tr><w:tc><w:tcPr>"
        '<w:gridSpan w:val="3"/></w:tcPr><w:p><w:r><w:t>spanned</w:t></w:r></w:p>'
        "</w:tc><w:tc><w:p><w:r><w:t>last</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
    )
    result = parse_document(_docx(tmp_path / "span3.docx", body))
    assert "spanned | last" in result.full_text
    assert result.full_text.count("spanned") == 1


def test_the_reviewers_vertical_merge_chain_parses_without_recursing(tmp_path: Path) -> None:
    first = (
        '<w:tr><w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr>'
        "<w:p><w:r><w:t>x</w:t></w:r></w:p></w:tc></w:tr>"
    )
    rest = "<w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc></w:tr>" * 5000
    result = bounded_parse(
        _docx(tmp_path / "v.docx", _PARAGRAPH + f"<w:tbl>{first}{rest}</w:tbl>")
    ).result
    assert "x" in result.full_text


def _row_of(columns: int) -> str:
    cells = "".join(f"<w:tc><w:p><w:r><w:t>c{i}</w:t></w:r></w:p></w:tc>" for i in range(columns))
    return _PARAGRAPH + f"<w:tbl><w:tr>{cells}</w:tr></w:tbl>"


def test_a_table_row_at_the_column_cap_is_read_whole(tmp_path: Path) -> None:
    columns = docx_parser.MAX_TABLE_COLUMNS
    result = parse_document(_docx(tmp_path / "c.docx", _row_of(columns)))
    assert FLAG_PARTIAL_EXTRACTION not in result.flags
    assert f"c{columns - 1}" in result.full_text


def test_a_table_row_one_column_past_the_cap_is_cut(tmp_path: Path) -> None:
    columns = docx_parser.MAX_TABLE_COLUMNS + 1
    result = parse_document(_docx(tmp_path / "c.docx", _row_of(columns)))
    assert FLAG_PARTIAL_EXTRACTION in result.flags
    assert f"c{columns - 1}" not in result.full_text


def _paragraphs(count: int) -> str:
    return "".join(f"<w:p><w:r><w:t>p{i}</w:t></w:r></w:p>" for i in range(count))


def test_a_document_at_the_block_budget_is_read_whole(tmp_path: Path) -> None:
    count = docx_parser.MAX_DOCX_BLOCKS
    result = parse_document(_docx(tmp_path / "b.docx", _paragraphs(count)))
    assert FLAG_PARTIAL_EXTRACTION not in result.flags
    assert f"p{count - 1}" in result.full_text


def test_a_document_one_block_past_the_budget_is_cut(tmp_path: Path) -> None:
    count = docx_parser.MAX_DOCX_BLOCKS + 1
    result = parse_document(_docx(tmp_path / "b.docx", _paragraphs(count)))
    assert FLAG_PARTIAL_EXTRACTION in result.flags
    assert f"p{count - 1}" not in result.full_text


def test_a_genuine_report_with_tables_still_parses(tmp_path: Path) -> None:
    """Valid input passes: a real python-docx report with headings, runs and tables."""
    document = docx.Document()
    for section in range(40):
        document.add_heading(f"Section {section}", level=1)
        for line in range(25):
            paragraph = document.add_paragraph()
            paragraph.add_run(f"Sentence {section}.{line} ").bold = True
            paragraph.add_run("continues here.")
        table = document.add_table(rows=10, cols=6)
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                cell.text = f"r{r}c{c}"
    path = tmp_path / "report.docx"
    document.save(path)
    result = parse_document(path)
    assert FLAG_PARTIAL_EXTRACTION not in result.flags
    assert "Sentence 39.24" in result.full_text
    assert "r9c5" in result.full_text


def test_a_memory_error_while_opening_a_docx_is_never_relabelled(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The open is inside the adapter's catch-all, so this is where relabelling could happen."""
    import docx as real_docx

    # Build the file BEFORE patching: the builder itself calls docx.Document, and a test
    # that builds after the patch passes on the builder's MemoryError, vacuously.
    path = _docx(tmp_path / "mo.docx", _PARAGRAPH)
    opened: list[str] = []

    def _boom(*args: object, **_kwargs: object) -> object:
        opened.append(str(args[0]) if args else "")
        raise MemoryError

    monkeypatch.setattr(real_docx, "Document", _boom)
    with pytest.raises(MemoryError):
        parse_document(path)
    assert opened == [str(path)], "precondition: the parser's own open raised it"


def test_a_memory_error_in_a_docx_is_never_relabelled(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _boom(*_args: object, **_kwargs: object) -> object:
        raise MemoryError

    monkeypatch.setattr(_live("docx"), "_extract_sections", _boom)
    with pytest.raises(MemoryError):
        parse_document(_docx(tmp_path / "m.docx", _PARAGRAPH))


def test_a_recursion_error_in_a_docx_is_a_clean_refusal(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _deep(*_args: object, **_kwargs: object) -> object:
        raise RecursionError

    monkeypatch.setattr(_live("docx"), "_extract_sections", _deep)
    with pytest.raises(CorruptDocumentError) as caught:
        parse_document(_docx(tmp_path / "r.docx", _PARAGRAPH))
    assert caught.value.context["reason"] == "too_complex"


# --------------------------------------------------------------------------- pdf


def test_a_pdf_declaring_more_pages_than_the_limit_is_refused_before_reading(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Refused on the claimed count, before the page tree is flattened."""
    import pypdf

    flattened: list[int] = []
    real_flatten = pypdf.PdfReader._flatten

    def _spy(self: object, *args: object, **kwargs: object) -> object:
        flattened.append(1)
        return real_flatten(self, *args, **kwargs)  # type: ignore[arg-type]

    monkeypatch.setattr(pypdf.PdfReader, "_flatten", _spy)
    path = _pdf(tmp_path / "claim.pdf", pages=1, declared=pdf_parser.MAX_PDF_PAGES + 1)
    with pytest.raises(CorruptDocumentError) as caught:
        parse_document(path)
    assert caught.value.context["reason"] == "too_many_pages"
    assert flattened == []


def test_a_pdf_at_the_page_limit_is_not_refused(tmp_path: Path) -> None:
    path = _pdf(tmp_path / "at.pdf", pages=pdf_parser.MAX_PDF_PAGES)
    result = parse_document(path)
    assert result.page_count == pdf_parser.MAX_PDF_PAGES


def test_a_pdf_lying_about_its_page_count_is_refused_on_the_real_count(tmp_path: Path) -> None:
    path = _pdf(tmp_path / "lie.pdf", pages=pdf_parser.MAX_PDF_PAGES + 1, declared=3)
    with pytest.raises(CorruptDocumentError) as caught:
        parse_document(path)
    assert caught.value.context["reason"] == "too_many_pages"


def test_text_is_extracted_from_every_page_at_the_text_limit(tmp_path: Path) -> None:
    path = _pdf(tmp_path / "t.pdf", pages=pdf_parser.MAX_PDF_TEXT_PAGES)
    result = parse_document(path)
    assert FLAG_PARTIAL_EXTRACTION not in result.flags
    assert len(result.sections) == pdf_parser.MAX_PDF_TEXT_PAGES


def test_text_stops_one_page_past_the_text_limit_and_is_flagged(tmp_path: Path) -> None:
    """Many pages sharing one stream: the work is per page, so pages are what is bounded."""
    path = _pdf(tmp_path / "t.pdf", pages=pdf_parser.MAX_PDF_TEXT_PAGES + 1)
    result = parse_document(path)
    assert FLAG_PARTIAL_EXTRACTION in result.flags
    assert len(result.sections) == pdf_parser.MAX_PDF_TEXT_PAGES
    assert result.page_count == pdf_parser.MAX_PDF_TEXT_PAGES + 1
