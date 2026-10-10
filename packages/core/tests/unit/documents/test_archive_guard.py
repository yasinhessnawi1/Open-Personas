"""The archive guard refuses hostile Office archives before a library opens them (D-FV0-08).

Each limit is crossed in its real unit: entries by count, sizes in bytes, and the lying
header by the bytes a parser would actually have to inflate. Archives that would be slow
to build for real (hundreds of megabytes declared) are built small and then have their
central directory rewritten, because the guard's first stage reads only what the
directory declares.
"""

from __future__ import annotations

import io
import struct
import tracemalloc
import zipfile
from pathlib import Path

import docx
import openpyxl
import pytest
from persona.documents import archive_guard
from persona.documents.archive_guard import (
    MAX_ARCHIVE_ENTRIES,
    MAX_ENTRY_UNCOMPRESSED_BYTES,
    MAX_TOTAL_UNCOMPRESSED_BYTES,
    check_ooxml_archive,
)
from persona.documents.errors import CorruptDocumentError
from persona.documents.parsers import (
    _DISPATCH_TABLE,
    _OOXML_FORMAT_BY_EXTENSION,
    SUPPORTED_EXTENSIONS,
    parse_document,
)

from tests.unit.documents._bounded_parse import bounded_parse

FIXTURE_DIR = Path(__file__).resolve().parents[2] / "fixtures" / "documents"

_CONTENT_TYPES = (
    b'<?xml version="1.0" encoding="UTF-8"?>'
    b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>'
)


# --------------------------------------------------------------------------- builders


def _docx_bytes(text: str = "hello") -> bytes:
    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "hello"  # type: ignore[index]
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _zip(entries: list[tuple[str, bytes]], *, method: int = zipfile.ZIP_DEFLATED) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", method) as archive:
        for name, data in entries:
            archive.writestr(name, data)
    return buffer.getvalue()


def _rewrite(entries: dict[str, bytes], replace: dict[str, bytes]) -> bytes:
    return _zip([(name, replace.get(name, data)) for name, data in entries.items()])


def _members(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {info.filename: archive.read(info.filename) for info in archive.infolist()}


def _forge_directory(
    data: bytes,
    name: str,
    *,
    file_size: int | None = None,
    flag_bits: int | None = None,
    crc: int | None = None,
    compress_size: int | None = None,
    header_offset: int | None = None,
) -> bytes:
    """Rewrite one entry's central directory record; the stored bytes are untouched."""
    raw = bytearray(data)
    end = raw.rfind(b"PK\x05\x06")
    (count,) = struct.unpack_from("<H", raw, end + 10)
    (offset,) = struct.unpack_from("<I", raw, end + 16)
    found = False
    for _ in range(count):
        assert raw[offset : offset + 4] == b"PK\x01\x02"
        name_length, extra_length, comment_length = struct.unpack_from("<HHH", raw, offset + 28)
        if raw[offset + 46 : offset + 46 + name_length].decode() == name:
            found = True
            if flag_bits is not None:
                struct.pack_into("<H", raw, offset + 8, flag_bits)
            if crc is not None:
                struct.pack_into("<I", raw, offset + 16, crc)
            if compress_size is not None:
                struct.pack_into("<I", raw, offset + 20, compress_size)
            if file_size is not None:
                struct.pack_into("<I", raw, offset + 24, file_size)
            if header_offset is not None:
                struct.pack_into("<I", raw, offset + 42, header_offset)
        offset += 46 + name_length + extra_length + comment_length
    assert found, f"no entry {name!r} to forge"
    return bytes(raw)


def _docx_with_media() -> bytes:
    """A real docx plus one binary part, for limits that apply to non-XML entries."""
    members = _members(_docx_bytes())
    members["word/media/image1.bin"] = b"\x89PNG" + b"\x00" * 2048
    return _zip(list(members.items()))


def _write(tmp_path: Path, name: str, data: bytes) -> Path:
    path = tmp_path / name
    path.write_bytes(data)
    return path


def _reason(path: Path, document_format: str = "docx") -> str:
    with pytest.raises(CorruptDocumentError) as caught:
        check_ooxml_archive(path, document_format=document_format)
    return caught.value.context["reason"]


# --------------------------------------------------------------------------- real files pass


@pytest.mark.parametrize(
    ("name", "document_format"), [("sample.docx", "docx"), ("sample.xlsx", "xlsx")]
)
def test_real_office_fixtures_pass_the_guard(name: str, document_format: str) -> None:
    check_ooxml_archive(FIXTURE_DIR / name, document_format=document_format)


def test_freshly_generated_docx_and_xlsx_pass_the_guard(tmp_path: Path) -> None:
    check_ooxml_archive(_write(tmp_path, "a.docx", _docx_bytes()), document_format="docx")
    check_ooxml_archive(_write(tmp_path, "a.xlsx", _xlsx_bytes()), document_format="xlsx")


def test_a_stored_entry_passes_the_guard(tmp_path: Path) -> None:
    data = _zip([("[Content_Types].xml", _CONTENT_TYPES), ("media/a.bin", b"x" * 5000)], method=0)
    check_ooxml_archive(_write(tmp_path, "s.docx", data), document_format="docx")


# --------------------------------------------------------------------------- valid input passes
#
# Every refusal below is tested hard; the property that a VALID archive passes has to be
# tested as hard, or a guard that refuses too much ships green. These are that half.

#: Sizes that a 37-of-400 false-refusal fuzz (orchestrator, seed 7) hit on the first
#: version of the guard: zlib still held output when the compressed input ran out inside
#: a full 64 KiB output step, and the guard stopped counting.
_KNOWN_FALSE_REFUSALS = (
    (393_450, "repeat"),
    (720_942, "xml"),
    (786_575, "repeat"),
    (1_245_268, "xml"),
    (1_769_484, "xml"),
    (2_293_819, "repeat"),
    (2_359_386, "repeat"),
)


def _payload(size: int, kind: str, seed: int) -> bytes:
    if kind == "repeat":
        return b"a" * size
    if kind == "xml":
        # Whole elements, then plain text to the exact size, so it stays well-formed XML.
        unit = b"<w:p><w:r><w:t>hello</w:t></w:r></w:p>"
        whole = size // len(unit)
        return unit * whole + b"x" * (size - whole * len(unit))
    # Two-symbol noise: compressible, but not trivially, like real XML runs.
    import random

    noise = random.Random(seed).randbytes(size)
    return noise.translate(bytes(97 + (b & 1) for b in range(256)))


def _valid_cases() -> list[tuple[int, str, int]]:
    cases = [(size, kind, level) for size, kind in _KNOWN_FALSE_REFUSALS for level in (1, 6, 9)]
    for multiple in (1, 2, 3, 6, 12, 17, 25, 38):
        for offset in (-300, -1, 0, 1, 300):
            size = 64 * 1024 * multiple + offset
            for index, kind in enumerate(("repeat", "xml", "noise")):
                level = (1, 6, 9)[(multiple + offset + index) % 3]
                if kind == "noise" and level == 9 and multiple > 3:
                    level = 6  # level 9 on large noise costs seconds per file
                cases.append((size, kind, level))
    return cases


def test_valid_deflate_archives_of_every_size_pass_the_guard(tmp_path: Path) -> None:
    """Deterministic port of the false-refusal fuzz: zero refusals allowed."""
    cases = _valid_cases()
    assert len(cases) > 100, "the corpus is too small to mean anything"
    refused: list[tuple[int, str, int, str]] = []
    for number, (size, kind, level) in enumerate(cases):
        path = tmp_path / f"v{number}.docx"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=level) as archive:
            archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
            body = _payload(size, kind, seed=number)
            archive.writestr("word/document.xml", b"<w:document>" + body + b"</w:document>")
        try:
            check_ooxml_archive(path, document_format="docx")
        except CorruptDocumentError as exc:
            refused.append((size, kind, level, exc.context["reason"]))
    assert refused == []


#: Valid parts whose deflate stream ends while zlib still holds output it could not hand
#: back inside the last step: all input is consumed, the stream is not at its end yet, and
#: only a drain step reads the rest. Found by a randomised search (T0 task 6), because the
#: corpus above never reached that state.
_ENDS_WITH_OUTPUT_BUFFERED = ((65_547, "ab", 6), (1_245_311, "ab", 9), (2_031_819, "xml", 1))


@pytest.mark.parametrize(("size", "kind", "level"), _ENDS_WITH_OUTPUT_BUFFERED)
def test_a_part_whose_stream_ends_with_output_still_buffered_passes(
    tmp_path: Path, size: int, kind: str, level: int
) -> None:
    if kind == "ab":
        body = (b"ab" * (size // 2 + 1))[:size]
    else:
        unit = b"<w:p><w:r><w:t>hello</w:t></w:r></w:p>"
        body = (unit * (size // len(unit) + 1))[:size]
    path = tmp_path / "buffered.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=level) as archive:
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        archive.writestr("word/media/blob.bin", body)
    check_ooxml_archive(path, document_format="docx")


def test_every_real_office_file_in_the_repo_passes_the_guard() -> None:
    """Corpus: every well-formed docx and xlsx the repository carries passes.

    The two ``corrupt.*`` fixtures exist to be unreadable and are excluded by name.
    """
    root = Path(__file__).resolve().parents[5]
    files = sorted(
        path
        for pattern in ("*.docx", "*.xlsx")
        for path in (root / "packages").rglob(pattern)
        if "node_modules" not in path.parts and not path.name.startswith("corrupt.")
    )
    assert len(files) >= 3, f"corpus too small to mean anything: {files}"
    for path in files:
        check_ooxml_archive(path, document_format=path.suffix.lstrip("."))


def test_a_deflate_stream_without_its_final_block_is_refused(tmp_path: Path) -> None:
    """The stream must reach its end, not merely produce the declared bytes.

    The entry is a raw deflate stream flushed with Z_SYNC_FLUSH and never finished: zlib
    hands back every byte and the CRC matches, but the final block never comes.
    """
    import zlib

    body = b"<w:document>" + b"x" * 10_000 + b"</w:document>"
    compressor = zlib.compressobj(6, zlib.DEFLATED, -zlib.MAX_WBITS)
    unfinished = compressor.compress(body) + compressor.flush(zlib.Z_SYNC_FLUSH)
    inflater = zlib.decompressobj(-zlib.MAX_WBITS)
    assert inflater.decompress(unfinished) == body
    assert not inflater.eof, "precondition: the stream really is unfinished"

    archive_bytes = _zip(
        [("[Content_Types].xml", _CONTENT_TYPES), ("word/document.xml", b"placeholder")],
        method=zipfile.ZIP_STORED,
    )
    raw = bytearray(archive_bytes)
    # Rebuild the document entry by hand: local header + the unfinished deflate bytes.
    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
        content_types = archive.getinfo("[Content_Types].xml")
        document = archive.getinfo("word/document.xml")
    name = b"word/document.xml"
    local = (
        struct.pack(
            "<4sHHHHHIIIHH",
            b"PK\x03\x04",
            20,
            0,
            zipfile.ZIP_DEFLATED,
            0,
            0,
            zlib.crc32(body),
            len(unfinished),
            len(body),
            len(name),
            0,
        )
        + name
        + unfinished
    )
    central = (
        struct.pack(
            "<4sHHHHHHIIIHHHHHII",
            b"PK\x01\x02",
            20,
            20,
            0,
            zipfile.ZIP_DEFLATED,
            0,
            0,
            zlib.crc32(body),
            len(unfinished),
            len(body),
            len(name),
            0,
            0,
            0,
            0,
            0,
            document.header_offset,
        )
        + name
    )
    head = bytes(raw[: document.header_offset])
    end = raw.rfind(b"PK\x05\x06")
    (cd_offset,) = struct.unpack_from("<I", raw, end + 16)
    first_central = bytes(raw[cd_offset : cd_offset + 46 + len(content_types.filename)])
    cd_start = len(head) + len(local)
    directory = first_central + central
    eocd = struct.pack("<4sHHHHIIH", b"PK\x05\x06", 0, 0, 2, 2, len(directory), cd_start, 0)
    data = head + local + directory + eocd

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert archive.read("word/document.xml") == body, "precondition: zipfile accepts it"
    assert _reason(_write(tmp_path, "unfinished.docx", data)) == "size_mismatch"


# --------------------------------------------------------------------------- shape refusals


def test_bytes_that_are_not_a_zip_are_refused(tmp_path: Path) -> None:
    assert _reason(_write(tmp_path, "x.docx", b"not an archive at all")) == "not_a_zip"


def test_data_prefixed_before_the_archive_is_refused(tmp_path: Path) -> None:
    # zipfile alone accepts this (it reads from the end); the polyglot shape is refused.
    polyglot = b"GIF89a" + b"\x00" * 64 + _docx_bytes()
    with zipfile.ZipFile(io.BytesIO(polyglot)) as archive:
        assert archive.namelist(), "precondition: zipfile accepts the prefixed archive"
    assert _reason(_write(tmp_path, "p.docx", polyglot)) == "not_a_zip"


def test_an_archive_at_the_entry_limit_passes(tmp_path: Path) -> None:
    entries = [("[Content_Types].xml", _CONTENT_TYPES)]
    entries += [(f"p/{i}.bin", b"") for i in range(MAX_ARCHIVE_ENTRIES - 1)]
    path = _write(tmp_path, "n.docx", _zip(entries))
    with zipfile.ZipFile(path) as archive:
        assert len(archive.infolist()) == MAX_ARCHIVE_ENTRIES
    check_ooxml_archive(path, document_format="docx")


def test_an_archive_one_entry_over_the_limit_is_refused(tmp_path: Path) -> None:
    entries = [("[Content_Types].xml", _CONTENT_TYPES)]
    entries += [(f"p/{i}.bin", b"") for i in range(MAX_ARCHIVE_ENTRIES)]
    path = _write(tmp_path, "n.docx", _zip(entries))
    with zipfile.ZipFile(path) as archive:
        assert len(archive.infolist()) == MAX_ARCHIVE_ENTRIES + 1
    assert _reason(path) == "too_many_entries"


def test_duplicate_entry_names_are_refused(tmp_path: Path) -> None:
    with pytest.warns(UserWarning, match="Duplicate name"):
        data = _zip([("[Content_Types].xml", _CONTENT_TYPES), ("a.xml", b"1"), ("a.xml", b"2")])
    assert _reason(_write(tmp_path, "d.docx", data)) == "duplicate_entry"


def test_an_encrypted_entry_is_refused(tmp_path: Path) -> None:
    data = _forge_directory(_docx_bytes(), "word/document.xml", flag_bits=0x1)
    assert _reason(_write(tmp_path, "e.docx", data)) == "encrypted_entry"


def test_an_archive_without_content_types_is_refused(tmp_path: Path) -> None:
    members = _members(_docx_bytes())
    members.pop("[Content_Types].xml")
    data = _zip(list(members.items()))
    assert _reason(_write(tmp_path, "m.docx", data)) == "missing_part"


def test_an_unsupported_compression_method_is_refused(tmp_path: Path) -> None:
    data = _zip([("[Content_Types].xml", _CONTENT_TYPES)], method=zipfile.ZIP_BZIP2)
    assert _reason(_write(tmp_path, "b.docx", data)) == "unsupported_compression"


# --------------------------------------------------------------------------- declared sizes


def test_an_entry_declaring_exactly_the_entry_limit_reaches_the_stream_check(
    tmp_path: Path,
) -> None:
    # Declaring the limit is allowed by the directory stage; the stream stage then sees
    # the real (small) content does not match the declaration. The reason proves which
    # stage refused: at the limit is not "entry_too_large".
    data = _forge_directory(
        _docx_with_media(), "word/media/image1.bin", file_size=MAX_ENTRY_UNCOMPRESSED_BYTES
    )
    assert _reason(_write(tmp_path, "l.docx", data)) == "size_mismatch"


def test_an_entry_declaring_one_byte_over_the_entry_limit_is_refused(tmp_path: Path) -> None:
    data = _forge_directory(
        _docx_with_media(), "word/media/image1.bin", file_size=MAX_ENTRY_UNCOMPRESSED_BYTES + 1
    )
    assert _reason(_write(tmp_path, "l.docx", data)) == "entry_too_large"


def _three_entries_declaring(total: int) -> bytes:
    data = _zip(
        [
            ("[Content_Types].xml", _CONTENT_TYPES),
            ("a.bin", b"a"),
            ("b.bin", b"b"),
        ]
    )
    first = MAX_ENTRY_UNCOMPRESSED_BYTES
    data = _forge_directory(data, "a.bin", file_size=first)
    content_types = len(_CONTENT_TYPES)
    return _forge_directory(data, "b.bin", file_size=total - first - content_types)


def test_an_archive_declaring_exactly_the_total_limit_reaches_the_stream_check(
    tmp_path: Path,
) -> None:
    path = _write(tmp_path, "t.docx", _three_entries_declaring(MAX_TOTAL_UNCOMPRESSED_BYTES))
    with zipfile.ZipFile(path) as archive:
        assert sum(i.file_size for i in archive.infolist()) == MAX_TOTAL_UNCOMPRESSED_BYTES
    assert _reason(path) == "size_mismatch"


def test_an_archive_declaring_one_byte_over_the_total_limit_is_refused(tmp_path: Path) -> None:
    path = _write(tmp_path, "t.docx", _three_entries_declaring(MAX_TOTAL_UNCOMPRESSED_BYTES + 1))
    with zipfile.ZipFile(path) as archive:
        assert sum(i.file_size for i in archive.infolist()) == MAX_TOTAL_UNCOMPRESSED_BYTES + 1
    assert _reason(path) == "archive_too_large"


# --------------------------------------------------------------------------- the lying header


_INFLATED_BYTES = 40 * 1024 * 1024
_DECLARED_BYTES = 4096


def _lying_docx() -> bytes:
    """A docx whose document.xml inflates to 40 MiB while declaring 4 KiB."""
    members = _members(_docx_bytes())
    padding = b"<!--" + b"A" * _INFLATED_BYTES + b"-->"
    real = members["word/document.xml"] + padding
    return _forge_directory(
        _rewrite(members, {"word/document.xml": real}),
        "word/document.xml",
        file_size=_DECLARED_BYTES,
    )


def test_the_lying_header_is_a_real_threat_without_the_guard(tmp_path: Path) -> None:
    """Precondition: the stdlib really inflates far past the declared size on a plain read.

    Without this the refusal below could be passing for a harmless file.
    """
    path = _write(tmp_path, "lie.docx", _lying_docx())
    with zipfile.ZipFile(path) as archive:
        info = archive.getinfo("word/document.xml")
        assert info.file_size == _DECLARED_BYTES
        assert info.compress_size < 1024 * 1024
        tracemalloc.start()
        try:
            with pytest.raises(zipfile.BadZipFile):
                archive.read("word/document.xml")
            _, peak = tracemalloc.get_traced_memory()
        finally:
            tracemalloc.stop()
    assert peak >= _INFLATED_BYTES, f"stdlib peaked at {peak} bytes"


def test_the_lying_header_is_refused_without_inflating_past_the_declared_size(
    tmp_path: Path,
) -> None:
    path = _write(tmp_path, "lie.docx", _lying_docx())
    tracemalloc.start()
    try:
        reason = _reason(path)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert reason == "size_mismatch"
    # One 64 KiB step of output plus bookkeeping, nowhere near the 40 MiB the stream holds.
    assert peak < 2 * 1024 * 1024, f"guard peaked at {peak} bytes"


def test_a_header_matching_the_truncated_crc_is_still_refused(tmp_path: Path) -> None:
    """The stream must END at the declared size, not merely checksum to it.

    An attacker can set the CRC of the first N inflated bytes and declare N, so a check
    that stops reading at N and compares CRCs passes while the library's own unbounded
    read still inflates the rest.
    """
    import zlib

    members = _members(_docx_bytes())
    head = members["word/document.xml"]
    real = head + b"<!--" + b"A" * (8 * 1024 * 1024) + b"-->"
    data = _rewrite(members, {"word/document.xml": real})
    data = _forge_directory(data, "word/document.xml", file_size=len(head), crc=zlib.crc32(head))
    assert _reason(_write(tmp_path, "crc.docx", data)) == "size_mismatch"


def test_inflating_stops_one_step_past_the_declared_size(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The work bound, in bytes: the guard never inflates a lying entry to its end.

    Memory alone cannot show this (each step is small either way), so the inflater is
    wrapped to count every byte it hands back.
    """
    import zlib

    produced: list[int] = []
    real_decompressobj = zlib.decompressobj

    class _Counting:
        def __init__(self, *args: int) -> None:
            self._inner = real_decompressobj(*args)

        def decompress(self, data: bytes, max_length: int = 0) -> bytes:
            out = self._inner.decompress(data, max_length)
            produced.append(len(out))
            return out

        @property
        def unconsumed_tail(self) -> bytes:
            return self._inner.unconsumed_tail

        @property
        def eof(self) -> bool:
            return self._inner.eof

    # Build the file BEFORE patching: zipfile shares the zlib module with the guard.
    path = _write(tmp_path, "lie.docx", _lying_docx())
    monkeypatch.setattr(archive_guard.zlib, "decompressobj", _Counting)
    assert _reason(path) == "size_mismatch"
    assert produced, "precondition: the counting inflater was used"
    # Everything the other entries legitimately hold, plus the lying entry's declared
    # 4 KiB, plus at most one 64 KiB step: nowhere near the 40 MiB the stream holds.
    honest = sum(
        info.file_size
        for info in zipfile.ZipFile(path).infolist()
        if info.filename != "word/document.xml"
    )
    assert sum(produced) <= honest + _DECLARED_BYTES + 64 * 1024


def test_a_header_offset_past_the_end_of_the_file_is_refused(tmp_path: Path) -> None:
    data = _docx_bytes()
    data = _forge_directory(data, "word/document.xml", header_offset=len(data) + 10)
    assert _reason(_write(tmp_path, "off.docx", data)) == "size_mismatch"


@pytest.mark.timeout(10)
def test_compressed_bytes_the_file_does_not_hold_are_refused(tmp_path: Path) -> None:
    data = _forge_directory(_docx_bytes(), "word/document.xml", compress_size=2**31)
    assert _reason(_write(tmp_path, "cs.docx", data)) == "size_mismatch"


def test_an_entry_shorter_than_declared_is_refused(tmp_path: Path) -> None:
    members = _members(_docx_bytes())
    declared = len(members["word/document.xml"]) + 1
    data = _forge_directory(_docx_bytes(), "word/document.xml", file_size=declared)
    assert _reason(_write(tmp_path, "short.docx", data)) == "size_mismatch"


def test_a_wrong_crc_is_refused(tmp_path: Path) -> None:
    data = _forge_directory(_docx_bytes(), "word/document.xml", crc=0)
    assert _reason(_write(tmp_path, "crc0.docx", data)) == "size_mismatch"


# --------------------------------------------------------------------------- parse_document wiring


def test_parse_document_refuses_a_hostile_docx_before_the_parser_runs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import persona.documents.parsers.docx as docx_parser

    called: list[Path] = []
    monkeypatch.setattr(docx_parser, "parse_docx", lambda p: called.append(p))
    path = _write(tmp_path, "lie.docx", _lying_docx())
    with pytest.raises(CorruptDocumentError) as caught:
        parse_document(path)
    assert caught.value.context["reason"] == "size_mismatch"
    assert called == []


def test_parse_document_refuses_a_hostile_xlsx(tmp_path: Path) -> None:
    data = _forge_directory(_xlsx_bytes(), "xl/workbook.xml", file_size=1)
    with pytest.raises(CorruptDocumentError) as caught:
        parse_document(_write(tmp_path, "bad.xlsx", data))
    assert caught.value.context["reason"] == "size_mismatch"
    assert caught.value.context["format"] == "xlsx"


def test_parse_document_still_parses_real_office_files() -> None:
    assert "twelve months" in parse_document(FIXTURE_DIR / "sample.docx").full_text
    assert parse_document(FIXTURE_DIR / "sample.xlsx").sheet_names


def test_every_zip_based_parser_in_the_dispatch_table_is_guarded() -> None:
    """Structure, not behaviour: a new archive format cannot join the table unguarded.

    Every parser module is classified here. A module added to the dispatch table that is
    in neither set fails this test until someone decides whether it reads archives.
    """
    archive_modules = {"persona.documents.parsers.docx", "persona.documents.parsers.xlsx"}
    plain_modules = {
        "persona.documents.parsers.text",
        "persona.documents.parsers.csv",
        "persona.documents.parsers.pdf",
    }
    modules = {module for module, _ in _DISPATCH_TABLE.values()}
    assert modules, "the dispatch table is empty; the guard check would be vacuous"
    assert modules <= archive_modules | plain_modules, modules - (archive_modules | plain_modules)
    archive_extensions = {
        ext for ext in SUPPORTED_EXTENSIONS if _DISPATCH_TABLE[ext][0] in archive_modules
    }
    assert archive_extensions == {".docx", ".xlsx"}
    assert set(_OOXML_FORMAT_BY_EXTENSION) == archive_extensions


# --------------------------------------------------------------------------- XML entity bombs


#: The most resident memory a child parsing an entity bomb may reach: the interpreter and
#: the parser libraries (about 60 MiB), plus what the parser holds before it refuses.
_ENTITY_PEAK_BYTES: int = 192 * 1024 * 1024


def _laughs(depth: int = 9) -> bytes:
    entities = b'<!ENTITY lol "lolololololololololol">' + b"".join(
        b'<!ENTITY lol%d "%s">' % (i, (b"&lol%d;" % (i - 1) if i > 1 else b"&lol;") * 10)
        for i in range(1, depth + 1)
    )
    return b'<?xml version="1.0"?><!DOCTYPE lolz [' + entities + b"]>"


def test_an_entity_bomb_in_a_docx_is_refused_without_expanding(tmp_path: Path) -> None:
    members = _members(_docx_bytes())
    body = members["word/document.xml"].split(b"?>", 1)[1].replace(b"hello", b"&lol9;")
    data = _rewrite(members, {"word/document.xml": _laughs() + body})
    path = _write(tmp_path, "lol.docx", data)
    # The guard refuses it outright: Office XML never carries a DTD.
    assert _reason(path) == "xml_doctype"
    # Defence in depth: the parser alone, without the guard, still does not expand it.
    # In a memory-limited child, so a regression that expands it fails this test only;
    # the peak is the whole process's resident memory, native lxml allocations included.
    with pytest.raises(CorruptDocumentError) as caught:
        bounded_parse(path, target="persona.documents.parsers.docx:parse_docx")
    peak = caught.value.peak_rss_bytes  # type: ignore[attr-defined]
    assert peak < _ENTITY_PEAK_BYTES, f"parse peaked at {peak} bytes"


def test_an_entity_bomb_in_an_xlsx_is_refused_without_expanding(tmp_path: Path) -> None:
    members = _members(_xlsx_bytes())
    sheet = members["xl/worksheets/sheet1.xml"]
    sheet = sheet.replace(b"<v>", b"<v>&lol9;", 1) if b"<v>" in sheet else sheet
    inline = (
        b'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&lol9;</t></is></c></row>'
        b"</sheetData>"
    )
    head, _, tail = sheet.partition(b"<sheetData>")
    _, _, tail = tail.partition(b"</sheetData>")
    bomb = _laughs() + head.split(b"?>", 1)[-1] + inline + tail
    data = _rewrite(members, {"xl/worksheets/sheet1.xml": bomb})
    path = _write(tmp_path, "lol.xlsx", data)
    assert _reason(path, "xlsx") == "xml_doctype"
    # The stdlib parser openpyxl reads sheets with stops this at expat's amplification
    # limit (defaults: an 8 MiB activation threshold and a factor of 100), so memory is
    # bounded by expat, not by us; the full expansion of ``&lol9;`` would be gigabytes.
    with pytest.raises(CorruptDocumentError) as caught:
        bounded_parse(path, target="persona.documents.parsers.xlsx:parse_xlsx")
    peak = caught.value.peak_rss_bytes  # type: ignore[attr-defined]
    assert peak < _ENTITY_PEAK_BYTES, f"parse peaked at {peak} bytes"


# --------------------------------------------------------------------------- XML stage


def _docx_with_document(xml: bytes) -> bytes:
    return _rewrite(_members(_docx_bytes()), {"word/document.xml": xml})


_W = b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _flood(elements: int) -> bytes:
    """A document.xml with exactly ``elements`` elements (root, body, then ``<w:p/>``)."""
    return (
        b'<?xml version="1.0" encoding="UTF-8"?><w:document '
        + _W
        + b"><w:body>"
        + b"<w:p/>" * (elements - 2)
        + b"</w:body></w:document>"
    )


def _other_elements(document: bytes) -> int:
    """Elements in every whole-loaded part of a fresh docx except word/document.xml."""
    import xml.parsers.expat

    total = 0
    for name, data in _members(document).items():
        if name == "word/document.xml" or not name.endswith((".xml", ".rels")):
            continue
        count = [0]
        parser = xml.parsers.expat.ParserCreate()
        parser.StartElementHandler = lambda *_a, c=count: c.__setitem__(0, c[0] + 1)
        parser.Parse(data, True)
        total += count[0]
    return total


def test_an_archive_at_the_element_budget_passes(tmp_path: Path) -> None:
    others = _other_elements(_docx_bytes())
    data = _docx_with_document(_flood(archive_guard.MAX_XML_ELEMENTS - others))
    check_ooxml_archive(_write(tmp_path, "at.docx", data), document_format="docx")


def test_an_archive_one_element_over_the_budget_is_refused(tmp_path: Path) -> None:
    others = _other_elements(_docx_bytes())
    data = _docx_with_document(_flood(archive_guard.MAX_XML_ELEMENTS - others + 1))
    assert _reason(_write(tmp_path, "over.docx", data)) == "too_many_elements"


def test_the_reviewers_paragraph_flood_is_refused(tmp_path: Path) -> None:
    """16 MiB of empty paragraphs: 2.8 million elements, 830 MiB in python-docx."""
    data = _docx_with_document(_flood(16 * 1024 * 1024 // len(b"<w:p/>")))
    assert _reason(_write(tmp_path, "flood.docx", data)) == "too_many_elements"


def test_a_whole_loaded_xml_part_at_the_byte_cap_reaches_the_stream_check(
    tmp_path: Path,
) -> None:
    data = _forge_directory(
        _docx_bytes(), "word/document.xml", file_size=archive_guard.MAX_XML_PART_BYTES
    )
    assert _reason(_write(tmp_path, "x.docx", data)) == "size_mismatch"


def test_a_whole_loaded_xml_part_one_byte_over_the_cap_is_refused(tmp_path: Path) -> None:
    data = _forge_directory(
        _docx_bytes(), "word/document.xml", file_size=archive_guard.MAX_XML_PART_BYTES + 1
    )
    assert _reason(_write(tmp_path, "x.docx", data)) == "xml_part_too_large"


def test_a_streamed_worksheet_may_exceed_the_xml_part_cap(tmp_path: Path) -> None:
    """Worksheets are read row by row, so only the entry cap applies to them."""
    data = _forge_directory(
        _xlsx_bytes(), "xl/worksheets/sheet1.xml", file_size=archive_guard.MAX_XML_PART_BYTES + 1
    )
    assert _reason(_write(tmp_path, "w.xlsx", data), "xlsx") == "size_mismatch"


def test_a_streamed_worksheet_does_not_spend_the_element_budget(tmp_path: Path) -> None:
    members = _members(_xlsx_bytes())
    rows = b"".join(
        b'<row r="%d"><c r="A%d"><v>%d</v></c></row>' % (i, i, i)
        for i in range(1, archive_guard.MAX_XML_ELEMENTS // 3 + 10)
    )
    sheet = (
        b'<?xml version="1.0" encoding="UTF-8"?>'
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b"<sheetData>" + rows + b"</sheetData></worksheet>"
    )
    data = _rewrite(members, {"xl/worksheets/sheet1.xml": sheet})
    check_ooxml_archive(_write(tmp_path, "big.xlsx", data), document_format="xlsx")


def test_a_doctype_in_any_xml_part_is_refused(tmp_path: Path) -> None:
    members = _members(_docx_bytes())
    styles = members["word/styles.xml"]
    head, sep, rest = styles.partition(b"?>")
    hostile = head + sep + b"<!DOCTYPE x>" + rest
    data = _rewrite(members, {"word/styles.xml": hostile})
    assert _reason(_write(tmp_path, "dt.docx", data)) == "xml_doctype"


def test_malformed_xml_is_refused(tmp_path: Path) -> None:
    data = _docx_with_document(b"<w:document " + _W + b"><w:body><w:p></w:body>")
    assert _reason(_write(tmp_path, "bad.docx", data)) == "malformed_xml"


def test_a_non_xml_part_is_not_scanned_as_xml(tmp_path: Path) -> None:
    check_ooxml_archive(_write(tmp_path, "m.docx", _docx_with_media()), document_format="docx")


def test_an_entry_name_with_control_characters_is_cleaned_in_the_refusal(
    tmp_path: Path,
) -> None:
    members = _members(_docx_bytes())
    members["word/evil\r\nFORGED: line\x1b[31m.xml"] = b"<a>"
    data = _zip(list(members.items()))
    with pytest.raises(CorruptDocumentError) as caught:
        check_ooxml_archive(_write(tmp_path, "c.docx", data), document_format="docx")
    entry = caught.value.context["entry"]
    assert caught.value.context["reason"] == "malformed_xml"
    assert entry == "word/evilFORGED: line[31m.xml"
    assert not any(ord(c) < 32 or ord(c) == 127 for c in entry)


def test_the_guard_module_imports_no_parser_library() -> None:
    """The guard is stdlib only, so it runs even where the extra is not installed."""
    source = Path(archive_guard.__file__).read_text(encoding="utf-8")
    for library in ("docx", "openpyxl", "pypdf", "pptx", "lxml"):
        assert f"import {library}" not in source
        assert f"from {library}" not in source


# --------------------------------------------------------------------------- content types (F6)


def _renamed_main(body_paragraphs: int, *, name: str = "word/document.bin") -> bytes:
    """The second review's ``binmain_*.docx``: the main part renamed, declared XML.

    ``[Content_Types].xml`` keeps declaring it as the main document, so python-docx
    loads it whole, but its name no longer ends in ``.xml``.
    """
    members = _members(_docx_bytes())
    main = members.pop("word/document.xml")
    head, _, tail = main.partition(b"<w:body>")
    main = head + b"<w:body>" + b"<w:p/>" * body_paragraphs + tail
    members[name] = main
    members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
        b"/word/document.xml", b"/" + name.encode()
    )
    members["_rels/.rels"] = members["_rels/.rels"].replace(b"word/document.xml", name.encode())
    return _zip(list(members.items()))


def test_a_renamed_main_part_still_spends_the_element_budget(tmp_path: Path) -> None:
    path = _write(tmp_path, "binmain.docx", _renamed_main(archive_guard.MAX_XML_ELEMENTS + 1))
    assert _reason(path) == "too_many_elements"


def test_a_renamed_main_part_over_the_xml_size_cap_is_refused(tmp_path: Path) -> None:
    paragraphs = archive_guard.MAX_XML_PART_BYTES // len(b"<w:p/>") + 1
    path = _write(tmp_path, "big.docx", _renamed_main(paragraphs))
    # The size cap comes first, from the directory entry, before it is inflated.
    assert _reason(path) == "xml_part_too_large"


def test_a_renamed_main_part_within_the_budgets_passes_and_parses(tmp_path: Path) -> None:
    """Valid input passes: a renamed main part is legal OPC, and python-docx reads it."""
    path = _write(tmp_path, "ok.docx", _renamed_main(1_000))
    check_ooxml_archive(path, document_format="docx")
    assert "hello" in parse_document(path).full_text


def test_a_renamed_main_part_with_a_doctype_is_refused(tmp_path: Path) -> None:
    members = _members(_renamed_main(1))
    members["word/document.bin"] = _laughs() + members["word/document.bin"].split(b"?>", 1)[1]
    path = _write(tmp_path, "lol.docx", _zip(list(members.items())))
    assert _reason(path) == "xml_doctype"


def test_binary_parts_are_not_scanned_as_xml(tmp_path: Path) -> None:
    """Valid input passes: an image, a printer-settings .bin and an SVG with a DOCTYPE
    are blobs to the libraries; scanning them as XML would refuse real files."""
    members = _members(_docx_bytes())
    svg = (
        b'<?xml version="1.0"?><!DOCTYPE svg PUBLIC "-//W3C//DTD SVG 1.1//EN" '
        b'"http://www.w3.org/Graphics/SVG/1.1/DTD/svg11.dtd"><svg/>'
    )
    members["word/media/image1.svg"] = svg
    members["word/media/image2.png"] = b"\x89PNG\r\n\x1a\n" + b"<" * 64
    members["word/printerSettings1.bin"] = b"\x00<not xml>\xff" * 32
    members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Default Extension="svg" ContentType="image/svg+xml"/>'
        b'<Default Extension="png" ContentType="image/png"/>'
        b'<Default Extension="bin" ContentType="application/vnd.openxmlformats-'
        b'officedocument.wordprocessingml.printerSettings"/></Types>',
    )
    path = _write(tmp_path, "media.docx", _zip(list(members.items())))
    check_ooxml_archive(path, document_format="docx")


def test_content_types_are_matched_without_regard_to_case(tmp_path: Path) -> None:
    members = _members(_renamed_main(archive_guard.MAX_XML_ELEMENTS + 1))
    members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
        b"/word/document.bin", b"/WORD/Document.BIN"
    )
    path = _write(tmp_path, "case.docx", _zip(list(members.items())))
    assert _reason(path) == "too_many_elements"


def test_a_content_types_part_over_the_xml_size_cap_is_refused_before_it_is_read(
    tmp_path: Path,
) -> None:
    """The content types part is read first, to classify the others; its own size is
    checked from the directory before that read."""
    members = _members(_docx_bytes())
    types = members["[Content_Types].xml"]
    padding = archive_guard.MAX_XML_PART_BYTES + 1 - len(types)
    members["[Content_Types].xml"] = types.replace(b"</Types>", b" " * padding + b"</Types>")
    assert len(members["[Content_Types].xml"]) == archive_guard.MAX_XML_PART_BYTES + 1
    path = _write(tmp_path, "types.docx", _zip(list(members.items())))
    assert _reason(path) == "xml_part_too_large"
