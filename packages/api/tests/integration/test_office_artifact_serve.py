"""R9-149: a generated spreadsheet, Word file or deck must open, not 404.

The reported symptom: the Files panel lists a generated ``.xlsx`` with a
spreadsheet icon, the user clicks it, and the fetch answers 404. Listing and
serving read two different maps. ``routes/artifacts.py`` has always listed the
three OOXML types with their real media types; ``image_service._media_type_for_ext``
did not know them, returned ``None``, and the serve route turned that into
``not_found`` for exactly the files Spec 24's document generation produces.

So this test refuses to assert on the map. For each of the three formats it:

1. builds REAL bytes with the same library the generation skill uses
   (openpyxl / python-docx / python-pptx) rather than a fake payload with an
   office extension;
2. persists them through the production :class:`WorkspaceDirPersister` (the
   adapter the chat-path tools call), so the ref and the F5 sidecar are the
   ones production writes;
3. reads the artifact listing the web's Files panel reads, and
4. fetches the very ref that listing returned, through
   ``GET /v1/personas/{id}/uploads/{ref:path}``.

Step 4 is the one that used to 404. It now returns the bytes, the OOXML media
type, ``nosniff``, and an attachment disposition (no browser renders these
inline, so a download is the only honest offer).
"""

from __future__ import annotations

import asyncio
import io
from pathlib import Path
from typing import TYPE_CHECKING

import pytest
from fastapi.testclient import TestClient
from persona_api.app import create_app
from persona_api.auth import AuthenticatedUser
from persona_api.config import APIConfig
from persona_api.middleware.rls_context import make_rls_engine
from persona_api.sandbox import (
    SandboxRequestContext,
    reset_sandbox_request_context,
    set_sandbox_request_context,
)
from persona_api.services.workspace_persister import WorkspaceDirPersister
from sqlalchemy import text

if TYPE_CHECKING:
    from collections.abc import Iterator

    from fastapi import FastAPI
    from sqlalchemy import Engine
    from tests.conftest import HashEmbedder384

pytestmark = pytest.mark.integration


_VALID_YAML = """\
schema_version: "1.0"
identity:
  name: Astrid
  role: Norwegian tenancy law assistant
  background: |
    Helps tenants understand husleieloven.
  language_default: en
  constraints: []
"""

_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx_bytes() -> bytes:
    """A real one-sheet workbook, written by openpyxl."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Month"
    sheet["B1"] = "Rent"
    sheet.append(["January", 12000])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    """A real one-paragraph document, written by python-docx."""
    from docx import Document

    document = Document()
    document.add_heading("Summary", level=1)
    document.add_paragraph("A tenancy summary the persona produced.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx_bytes() -> bytes:
    """A real one-slide deck, written by python-pptx."""
    from pptx import Presentation

    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


#: (media type, suggested filename, byte builder) per OOXML format.
_FORMATS: list[tuple[str, str, str]] = [
    (_XLSX, "quarterly-report.xlsx", "xlsx"),
    (_DOCX, "tenancy-summary.docx", "docx"),
    (_PPTX, "board-deck.pptx", "pptx"),
]

_BUILDERS = {"xlsx": _xlsx_bytes, "docx": _docx_bytes, "pptx": _pptx_bytes}


@pytest.fixture
def client(
    migrated_engine: Engine,  # noqa: ARG001 (ensures schema + grants)
    embedder: HashEmbedder384,
    tmp_path: Path,
) -> Iterator[tuple[TestClient, FastAPI, str]]:
    """Real FastAPI client + one user whose persona owns the artifacts."""
    import os

    app_url = os.environ.get("APP_DATABASE_URL")
    if not app_url:
        pytest.skip("APP_DATABASE_URL not set")

    cfg = APIConfig(
        app_database_url=app_url,
        audit_root=str(tmp_path / "audit"),
        workspace_root=tmp_path / "workspace",
    )
    app = create_app(cfg)

    async def _fake_verify(token: str) -> AuthenticatedUser:
        return AuthenticatedUser(id=token, email=None)

    user = "user_r9149"
    with TestClient(app) as c:
        app.state.verify_token = _fake_verify
        app.state.embedder = embedder
        # Same reason as test_charts_serve.py: the persona-detail capabilities
        # surface would lazily build a real chat backend without an API key.
        if hasattr(app.state, "tier_registry"):
            app.state.tier_registry = None
        su = make_rls_engine(os.environ["DATABASE_URL"])
        with su.begin() as conn:
            conn.execute(
                text("INSERT INTO users (id, email) VALUES (:i, :e) ON CONFLICT DO NOTHING"),
                {"i": user, "e": f"{user}@x.test"},
            )
        su.dispose()
        yield c, app, user
        su = make_rls_engine(os.environ["DATABASE_URL"])
        with su.begin() as conn:
            conn.execute(text("DELETE FROM users WHERE id = :i"), {"i": user})
        su.dispose()


def _auth(user_id: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {user_id}"}


def _create_persona(c: TestClient, user_id: str) -> str:
    resp = c.post("/v1/personas", json={"yaml": _VALID_YAML}, headers=_auth(user_id))
    assert resp.status_code == 201, resp.text
    return str(resp.json()["id"])


def _persist_generated(
    *,
    app: FastAPI,
    owner_id: str,
    persona_id: str,
    data: bytes,
    media_type: str,
    filename: str,
) -> str:
    """Persist ``data`` the way a generated artifact is persisted in production.

    Runs the real :class:`WorkspaceDirPersister` under a bound sandbox request
    context, so the workspace ref, the content addressing and the F5 sidecar all
    come from the shipping adapter rather than from this test.

    Returns:
        The workspace-relative ref the persister assigned.
    """

    async def _run() -> str:
        token = set_sandbox_request_context(
            SandboxRequestContext(owner_id=owner_id, conversation_id="conv_r9149")
        )
        try:
            persister = WorkspaceDirPersister(
                file_storage=app.state.file_storage, persona_id=persona_id
            )
            artifact = await persister.persist(
                data, mime_type=media_type, suggested_filename=filename
            )
            return artifact.workspace_path
        finally:
            reset_sandbox_request_context(token)

    return asyncio.run(_run())


def _listed_item(c: TestClient, user_id: str, persona_id: str, ref: str) -> dict[str, object]:
    resp = c.get(f"/v1/personas/{persona_id}/artifacts", headers=_auth(user_id))
    assert resp.status_code == 200, resp.text
    items = [item for item in resp.json()["items"] if item["ref"] == ref]
    assert items, f"{ref} not listed: {resp.json()}"
    return dict(items[0])


class TestGeneratedOfficeDocumentsAreServable:
    """The file the Files panel lists is the file the Files panel can open."""

    @pytest.mark.parametrize(("media_type", "filename", "builder"), _FORMATS)
    def test_listed_then_fetched_through_the_uploads_route(
        self,
        client: tuple[TestClient, FastAPI, str],
        media_type: str,
        filename: str,
        builder: str,
    ) -> None:
        c, app, user = client
        persona_id = _create_persona(c, user)
        data = _BUILDERS[builder]()

        ref = _persist_generated(
            app=app,
            owner_id=user,
            persona_id=persona_id,
            data=data,
            media_type=media_type,
            filename=filename,
        )

        # The panel lists it with the OOXML type (this half always worked).
        listed = _listed_item(c, user, persona_id, ref)
        assert listed["media_type"] == media_type

        # THE regression: opening that exact ref answered 404.
        resp = c.get(f"/v1/personas/{persona_id}/uploads/{ref}", headers=_auth(user))
        assert resp.status_code == 200, resp.text
        assert resp.content == data
        assert resp.headers["content-type"].startswith(media_type)
        assert resp.headers["x-content-type-options"] == "nosniff"
        assert resp.headers["content-disposition"].startswith("attachment")

    def test_a_sandbox_produced_name_serves_too(
        self,
        client: tuple[TestClient, FastAPI, str],
        tmp_path: Path,
    ) -> None:
        """Spec 16's document generation keeps the model's own filename.

        Its produced-file persister copies to ``uploads/<name>.<ext>`` rather
        than content-addressing it (D-F4-X-bare-ref-resolution), so the served
        ref carries a human name. Same map, same route, same contract.
        """
        c, _app, user = client
        persona_id = _create_persona(c, user)
        data = _xlsx_bytes()
        produced = tmp_path / "workspace" / user / persona_id / "uploads" / "rent-history.xlsx"
        produced.parent.mkdir(parents=True, exist_ok=True)
        produced.write_bytes(data)

        resp = c.get(
            f"/v1/personas/{persona_id}/uploads/uploads/rent-history.xlsx",
            headers=_auth(user),
        )
        assert resp.status_code == 200, resp.text
        assert resp.content == data
        assert resp.headers["content-type"].startswith(_XLSX)
        assert 'filename="rent-history.xlsx"' in resp.headers["content-disposition"]

    def test_a_non_ascii_filename_still_downloads(
        self,
        client: tuple[TestClient, FastAPI, str],
        tmp_path: Path,
    ) -> None:
        """A model names its own files, and the name is not always ASCII.

        Header values are encoded latin-1, so putting the raw name in
        ``filename=`` would turn a working download into a 500 the first time a
        persona wrote a Japanese or Greek filename. The quoted parameter stays
        ASCII and the real name rides along in ``filename*``.
        """
        c, _app, user = client
        persona_id = _create_persona(c, user)
        data = _xlsx_bytes()
        name = "kvartal_日本.xlsx"
        produced = tmp_path / "workspace" / user / persona_id / "uploads" / name
        produced.parent.mkdir(parents=True, exist_ok=True)
        produced.write_bytes(data)

        resp = c.get(
            f"/v1/personas/{persona_id}/uploads/uploads/{name}",
            headers=_auth(user),
        )
        assert resp.status_code == 200, resp.text
        assert resp.content == data
        disposition = resp.headers["content-disposition"]
        assert 'filename="kvartal_.xlsx"' in disposition
        assert "filename*=UTF-8''kvartal_%E6%97%A5%E6%9C%AC.xlsx" in disposition
